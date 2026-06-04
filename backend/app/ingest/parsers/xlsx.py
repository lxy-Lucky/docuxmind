"""Excel parser. One outline entry per sheet; segments grouped by row blocks."""
from datetime import date, datetime, time
from typing import Any

import openpyxl

BLOCK = 40


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def parse(path: str) -> tuple[dict, list[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        outline: dict[str, Any] = {"kind": "xlsx", "sheets": []}
        segments: list[dict] = []

        for sname in wb.sheetnames:
            ws = wb[sname]
            sheet_rows: list[tuple[int, str]] = []
            header: list[str] = []

            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                text = "\t".join(_cell_to_str(c) for c in row)
                if not text.strip():
                    continue
                if not header:
                    header = [_cell_to_str(c) for c in row][:8]
                sheet_rows.append((i, text))

            outline["sheets"].append(
                {"name": sname, "header": header, "row_count": len(sheet_rows)}
            )

            for k in range(0, len(sheet_rows), BLOCK):
                slc = sheet_rows[k : k + BLOCK]
                if not slc:
                    continue
                first, last = slc[0][0], slc[-1][0]
                content = "\n".join(f"行{i}: {t}" for i, t in slc)
                segments.append(
                    {
                        "locator": f"sheet={sname}|rows={first}-{last}",
                        "content": content,
                    }
                )

        return outline, segments
    finally:
        wb.close()
